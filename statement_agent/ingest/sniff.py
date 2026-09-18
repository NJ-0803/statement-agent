"""Structural discovery for tabular sources: encoding, delimiter, sheets, header row, candidate tables.

A CSV isn't self-describing, and a bank export rarely starts with its header: a bank name, account
details, a period line and blank rows usually come first. Neither assumption the old parsers made
("the first row is the header", "the active sheet is the right one") holds for real exports. This
module returns *ranked* candidates with human-readable evidence — never one opaque guess — and
mapping.py / the preview screen decide what to do with them.
"""

from __future__ import annotations

import codecs
import csv
import datetime as _dt
import io
import os
from collections import OrderedDict
from dataclasses import dataclass, field

from .vocab import header_role_scores, looks_like_amount, looks_like_date

HEADER_SCAN_ROWS = 50
MAX_TABLE_ROWS = 200_000
_DELIMITERS = [",", ";", "\t", "|"]
_DELIMITER_NAMES = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}


class TableTooLarge(ValueError):
    pass


@dataclass
class TableCandidate:
    sheet: str | None
    header_row: int | None  # 1-based source row of the header; None when no header row was found
    headers: list[str]  # de-duplicated; "Column N" where the source had no name
    rows: list[tuple[int, list[str]]]  # (1-based source row, cells) for every row after the header
    preamble: list[tuple[int, list[str]]] = field(default_factory=list)  # rows before the header
    score: float = 0.0
    evidence: list[str] = field(default_factory=list)
    duplicate_headers: list[str] = field(default_factory=list)
    row_pages: dict[int, int] = field(default_factory=dict)  # PDF tables: source row -> page number
    document_hints: dict = field(default_factory=dict)  # PDF tables: period, account, doc type, text outside the table

    @property
    def label(self) -> str:
        where = f"sheet '{self.sheet}', " if self.sheet else ""
        return f"{where}header on row {self.header_row}" if self.header_row else f"{where}no header row"


@dataclass
class SniffResult:
    kind: str  # "csv" | "xlsx"
    candidates: list[TableCandidate]  # best first
    encoding: str | None = None
    delimiter: str | None = None
    sheets: list[dict] = field(default_factory=list)  # {"name", "state", "rows"}
    formula_gaps: list[str] = field(default_factory=list)  # cells holding a formula whose result was never saved
    evidence: list[str] = field(default_factory=list)

    @property
    def best(self) -> TableCandidate | None:
        return self.candidates[0] if self.candidates else None

    def select(self, sheet: str | None, header_row: int | None) -> TableCandidate | None:
        for c in self.candidates:
            if c.sheet == sheet and c.header_row == header_row:
                return c
        return None


# ---------------------------------------------------------------------------
# Reading raw grids
# ---------------------------------------------------------------------------

def detect_encoding(data: bytes) -> tuple[str, str]:
    if data.startswith(codecs.BOM_UTF8):
        return "utf-8-sig", "UTF-8 byte-order mark"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16", "UTF-16 byte-order mark"
    head = data[:4000]
    # UTF-16 without a BOM is still valid UTF-8 byte-wise (NUL is legal UTF-8), so it has to be
    # recognised by its NUL pattern before a UTF-8 decode is even attempted.
    if len(head) >= 4:
        odd_nuls, even_nuls = head[1::2].count(0), head[0::2].count(0)
        half = len(head) // 2
        if odd_nuls > half * 0.4:
            return "utf-16-le", "NUL-byte pattern of UTF-16 (little-endian) text"
        if even_nuls > half * 0.4:
            return "utf-16-be", "NUL-byte pattern of UTF-16 (big-endian) text"
    try:
        data.decode("utf-8")
        return "utf-8", "decodes cleanly as UTF-8"
    except UnicodeDecodeError:
        pass
    try:
        data.decode("cp1252")
        return "cp1252", "not valid UTF-8; decoded as Windows-1252"
    except UnicodeDecodeError:
        return "latin-1", "not valid UTF-8 or Windows-1252; decoded as Latin-1"


def detect_delimiter(text: str) -> tuple[str, str]:
    sample = "\n".join(text.splitlines()[:60])
    best, best_score = ",", -1.0
    for d in _DELIMITERS:
        try:
            widths = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d) if any(c.strip() for c in r)]
        except csv.Error:
            continue
        if not widths:
            continue
        mode = max(set(widths), key=widths.count)
        if mode < 2:
            continue
        score = widths.count(mode) + mode * 0.01  # most rows agreeing on one width wins; wider breaks ties
        if score > best_score:
            best, best_score = d, score
    return best, f"{_DELIMITER_NAMES[best]}-separated (most rows share the same column count)"


def read_csv_grid(path: str) -> tuple[list[list[str]], str, str, list[str]]:
    with open(path, "rb") as f:
        data = f.read()
    encoding, enc_evidence = detect_encoding(data)
    text = data.decode(encoding, errors="replace")
    delimiter, delim_evidence = detect_delimiter(text)
    grid: list[list[str]] = []
    for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):  # handles quoted newlines
        grid.append([c.strip() for c in row])
        if len(grid) > MAX_TABLE_ROWS:
            raise TableTooLarge(f"more than {MAX_TABLE_ROWS:,} rows")
    return grid, encoding, delimiter, [f"Encoding: {enc_evidence}", f"Layout: {delim_evidence}"]


def cell_to_str(value) -> str:
    """Excel stores dates/numbers as typed values; a naive str() would turn a date cell into
    "2025-06-21 00:00:00", which date parsing rejects."""
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        return value.strftime("%Y-%m-%d") if value.time() == _dt.time() else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, _dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return str(value)
    return str(value).strip()


def find_uncached_formulas(path: str, limit: int = 20) -> list[str]:
    """Cells whose formula has no saved result. Excel and LibreOffice normally store the last computed value
    next to the formula, but a file written by a script (openpyxl, some bank portals) often has none — and
    then every such cell reads as empty. Those cells are named rather than silently dropped."""
    import openpyxl

    try:
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 - detection must never be the thing that fails an import
        return []
    gaps: list[str] = []
    try:
        for sheet in formulas.sheetnames:
            if sheet not in values.sheetnames:
                continue
            for formula_row, value_row in zip(formulas[sheet].iter_rows(), values[sheet].iter_rows()):
                for formula_cell, value_cell in zip(formula_row, value_row):
                    if value_cell.value in (None, "") and isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                        gaps.append(f"{sheet}!{formula_cell.coordinate} ({formula_cell.value})")
                        if len(gaps) >= limit:
                            return gaps
    finally:
        values.close()
        formulas.close()
    return gaps


def read_xlsx_grids(path: str) -> list[tuple[str, str, list[list[str]]]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            grid: list[list[str]] = []
            for raw in ws.iter_rows(values_only=True):
                grid.append([cell_to_str(v) for v in (raw or ())])
                if len(grid) > MAX_TABLE_ROWS:
                    raise TableTooLarge(f"sheet '{ws.title}' has more than {MAX_TABLE_ROWS:,} rows")
            sheets.append((ws.title, ws.sheet_state, grid))
        return sheets
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Finding tables inside a grid
# ---------------------------------------------------------------------------

def _trim(cells: list[str]) -> list[str]:
    end = len(cells)
    while end and not cells[end - 1]:
        end -= 1
    return cells[:end]


def _is_data_like(cells: list[str]) -> bool:
    return any(looks_like_date(c) for c in cells) and any(looks_like_amount(c) and not looks_like_date(c) for c in cells)


def find_tables(grid: list[list[str]], *, sheet: str | None = None) -> list[TableCandidate]:
    rows = [_trim(r) for r in grid]
    candidates: list[TableCandidate] = []
    seen_headers: set[tuple[str, ...]] = set()

    for i, cells in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in cells if c]
        if len(filled) < 2:
            continue
        signature = tuple(c.strip().lower() for c in cells)
        if signature in seen_headers:
            continue  # the same header repeated (page breaks in an export) belongs to the first table
        seen_headers.add(signature)
        if any(looks_like_date(c) or (looks_like_amount(c) and "." in c) for c in filled):
            continue  # header rows name columns; they don't hold dates or money
        hits = sum(1 for c in filled if max(header_role_scores(c).values(), default=0) >= 0.8)
        if hits == 0:
            continue
        following = [r for r in rows[i + 1:i + 11] if any(r)]
        follow = sum(1 for r in following if _is_data_like(r)) / len(following) if following else 0.0
        score = hits / len(filled) + min(hits, 4) * 0.5 + follow * 2
        evidence = [f"row {i + 1} names {hits} recognisable column(s)"]
        if follow:
            evidence.append(f"{follow:.0%} of the next rows look like transactions")
        else:
            evidence.append("no transaction-like rows directly below it")
        candidates.append(_build(rows, i, sheet, score, evidence))

    if not candidates:
        first_data = next((i for i, r in enumerate(rows) if _is_data_like(r)), None)
        if first_data is not None:
            width = max(len(r) for r in rows[first_data:]) if rows[first_data:] else 0
            cand = TableCandidate(
                sheet=sheet, header_row=None, headers=[f"Column {j + 1}" for j in range(width)],
                rows=[(n + 1, r) for n, r in enumerate(rows) if n >= first_data],
                preamble=[(n + 1, r) for n, r in enumerate(rows[:first_data]) if any(r)],
                score=0.1, evidence=["no header row found — columns can only be judged by their values"],
            )
            candidates.append(cand)

    candidates.sort(key=lambda c: -c.score)
    return candidates


def _build(rows: list[list[str]], header_idx: int, sheet: str | None, score: float, evidence: list[str]) -> TableCandidate:
    body = rows[header_idx + 1:]
    width = max([len(rows[header_idx])] + [len(r) for r in body]) if body else len(rows[header_idx])
    raw_headers = rows[header_idx] + [""] * (width - len(rows[header_idx]))
    headers, seen, dups = [], {}, []
    for j, h in enumerate(raw_headers):
        name = h or f"Column {j + 1}"
        if name.lower() in seen:
            seen[name.lower()] += 1
            dups.append(name)
            name = f"{name} ({seen[name.lower()]})"
        else:
            seen[name.lower()] = 1
        headers.append(name)
    if dups:
        evidence.append(f"duplicate column name(s) renamed: {sorted(set(dups))}")
    return TableCandidate(
        sheet=sheet, header_row=header_idx + 1, headers=headers,
        rows=[(header_idx + 2 + n, r) for n, r in enumerate(body)],
        preamble=[(n + 1, r) for n, r in enumerate(rows[:header_idx]) if any(r)],
        score=score, evidence=evidence, duplicate_headers=sorted(set(dups)),
    )


def sniff_csv(path: str) -> SniffResult:
    grid, encoding, delimiter, evidence = read_csv_grid(path)
    return SniffResult(kind="csv", candidates=find_tables(grid), encoding=encoding, delimiter=delimiter, evidence=evidence)


def sniff_xlsx(path: str) -> SniffResult:
    result = sniff_grids("xlsx", read_xlsx_grids(path))
    result.formula_gaps = find_uncached_formulas(path)
    if result.formula_gaps:
        result.evidence.append(f"{len(result.formula_gaps)} cell(s) hold a formula whose result was never saved")
    return result


def sniff_grids(kind: str, grids) -> SniffResult:
    candidates: list[TableCandidate] = []
    sheet_info = []
    for name, state, grid in grids:
        sheet_info.append({"name": name, "state": state, "rows": len(grid)})
        for cand in find_tables(grid, sheet=name):
            if state != "visible":
                cand.score *= 0.5
                cand.evidence.append(f"sheet is {state} in the workbook")
            candidates.append(cand)
    candidates.sort(key=lambda c: -c.score)
    if kind == "xlsx" or len(sheet_info) > 1:
        evidence = [f"File has {len(sheet_info)} sheet(s) or table(s): {', '.join(s['name'] for s in sheet_info)}"]
    else:
        evidence = [f"Read as a {kind.upper()} file"]
    return SniffResult(kind=kind, candidates=candidates, sheets=sheet_info, evidence=evidence)


_CACHE: "OrderedDict[tuple, SniffResult]" = OrderedDict()
_CACHE_SIZE = 16


def _cache_key(path: str):
    st = os.stat(path)
    return (os.path.realpath(path), st.st_size, st.st_mtime_ns)


def sniff_file(path: str) -> SniffResult:
    """Cached on (path, size, modified time): one import reads the same file several times (analyse, remap,
    review, commit), and re-reading a PDF's pages each time is the single most expensive part of an import."""
    try:
        key = _cache_key(path)
    except OSError:
        return _sniff_file(path)
    hit = _CACHE.get(key)
    if hit is not None:
        _CACHE.move_to_end(key)
        return hit
    result = _sniff_file(path)
    _CACHE[key] = result
    _CACHE.move_to_end(key)
    while len(_CACHE) > _CACHE_SIZE:
        _CACHE.popitem(last=False)
    return result


def _sniff_file(path: str) -> SniffResult:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return sniff_xlsx(path)
    if ext == ".pdf":
        from .pdf_columns import read_pdf_table

        table = read_pdf_table(path)
        return SniffResult(kind="pdf", candidates=[table] if table else [],
                           evidence=["Read the PDF's table by column position"] if table else ["No table header found in this PDF"])
    from .formats import TEXT_TABLE_EXTENSIONS, read_sheets

    if ext in (".csv", *TEXT_TABLE_EXTENSIONS) or not ext:
        return sniff_csv(path)
    return sniff_grids(ext.lstrip("."), read_sheets(path, ext))
